"""File -> table. CSV (UTF-8/Windows-1252, , ; tab |) and XLSX. Hard limits protect the server; nothing is executed or stored."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime

MAX_COLS = 60


class UploadError(ValueError):
    """Whole-file problems (unreadable, too big, wrong type). `code` is stable for clients."""
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code, self.message = code, message


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    lines: list[int]  # 1-based source line (CSV) or row number (XLSX) of each data row; the header is line 1
    encoding: str
    delimiter: str | None
    sheet: str | None = None


def _decode(data: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def _read_csv(data: bytes, max_rows: int) -> Table:
    text, enc = _decode(data)
    sample = text[:8192]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample.splitlines() else ""
        delim = max(",;\t|", key=first.count)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise UploadError("empty", "o arquivo está vazio") from None
    if len(headers) > MAX_COLS:
        raise UploadError("too_many_columns", f"colunas demais (máximo {MAX_COLS})")
    rows: list[list[str]] = []
    lines: list[int] = []
    for row in reader:
        if not any(c.strip() for c in row):
            continue
        if len(rows) >= max_rows:
            raise UploadError("too_many_rows", f"linhas demais (máximo {max_rows})")
        rows.append([c.strip() for c in row] + [""] * max(0, len(headers) - len(row)))
        lines.append(reader.line_num)
    return Table(headers, rows, lines, enc, delim)


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if (v.hour or v.minute or v.second) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_xlsx(data: bytes, max_rows: int) -> Table:
    from openpyxl import load_workbook  # local import: only needed for .xlsx
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise UploadError("bad_xlsx", "não consegui abrir a planilha (arquivo .xlsx inválido ou protegido)") from None
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        headers = [_cell(c) for c in next(it)]
    except StopIteration:
        raise UploadError("empty", "a planilha está vazia") from None
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) > MAX_COLS:
        raise UploadError("too_many_columns", f"colunas demais (máximo {MAX_COLS})")
    rows: list[list[str]] = []
    lines: list[int] = []
    for i, row in enumerate(it, start=2):
        cells = [_cell(c) for c in row[:len(headers)]]
        if not any(cells):
            continue
        if len(rows) >= max_rows:
            raise UploadError("too_many_rows", f"linhas demais (máximo {max_rows})")
        rows.append(cells + [""] * (len(headers) - len(cells)))
        lines.append(i)
    return Table(headers, rows, lines, "xlsx", None, ws.title)


def read_table(data: bytes, filename: str, *, max_bytes: int = 10_000_000, max_rows: int = 100_000) -> Table:
    if len(data) > max_bytes:
        raise UploadError("too_big", f"arquivo grande demais (máximo {max_bytes // 1_000_000} MB)")
    if not data.strip():
        raise UploadError("empty", "o arquivo está vazio")
    if data[:4] == b"PK\x03\x04":
        return _read_xlsx(data, max_rows)
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise UploadError("legacy_xls", "arquivos .xls antigos não são aceitos: salve como .xlsx ou .csv")
    if b"\x00" in data[:8192]:
        raise UploadError("binary", "o arquivo não parece ser texto (CSV) nem planilha (XLSX)")
    if not filename.lower().endswith((".csv", ".txt", ".tsv", ".xlsx")):
        raise UploadError("bad_type", "envie um arquivo .csv ou .xlsx")
    return _read_csv(data, max_rows)
